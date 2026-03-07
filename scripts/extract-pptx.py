#!/usr/bin/env python3
"""Extract content from PowerPoint (.pptx) files for slide-sage conversion.

Usage:
    python extract-pptx.py input.pptx [--output extracted.json]

Extracts slide content (text, images, notes, layout hints) into a structured
JSON format that the AI can use to generate an HTML presentation.

Requirements:
    pip install python-pptx Pillow
"""

from __future__ import annotations

import argparse
import base64
import json
import sys
from io import BytesIO
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("Error: python-pptx is required. Install with: pip install python-pptx")
    sys.exit(1)

try:
    from PIL import Image
except ImportError:
    Image = None


def extract_text_frame(text_frame) -> list[dict]:
    """Extract paragraphs from a text frame."""
    paragraphs = []
    for para in text_frame.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        alignment = "left"
        if para.alignment == PP_ALIGN.CENTER:
            alignment = "center"
        elif para.alignment == PP_ALIGN.RIGHT:
            alignment = "right"

        level = para.level or 0
        font_size = None
        is_bold = False

        for run in para.runs:
            if run.font.size:
                font_size = run.font.size.pt
            if run.font.bold:
                is_bold = True

        paragraphs.append({
            "text": text,
            "level": level,
            "alignment": alignment,
            "font_size": font_size,
            "bold": is_bold,
        })
    return paragraphs


def classify_element(paragraphs: list[dict]) -> str:
    """Classify text content as heading, body, or bullet."""
    if not paragraphs:
        return "empty"
    first = paragraphs[0]
    if first.get("bold") and first.get("font_size", 0) and first["font_size"] >= 24:
        return "heading"
    if any(p["level"] > 0 for p in paragraphs):
        return "bullets"
    return "body"


def extract_image(shape, output_dir: Path | None, image_id: int = 0) -> dict | None:
    """Extract image from a shape."""
    try:
        image = shape.image
        content_type = image.content_type
        blob = image.blob
        ext = content_type.split("/")[-1] if content_type else "png"
        if ext == "jpeg":
            ext = "jpg"

        result = {
            "type": "image",
            "content_type": content_type,
            "width": shape.width,
            "height": shape.height,
        }

        if output_dir:
            img_name = f"image_{image_id:03d}.{ext}"
            img_path = output_dir / img_name
            img_path.write_bytes(blob)
            result["path"] = str(img_path)
        else:
            if len(blob) < 37_500:
                encoded = base64.b64encode(blob).decode("ascii")
                result["data_uri"] = f"data:{content_type};base64,{encoded}"
            else:
                result["note"] = f"Image too large for inline ({len(blob)} bytes)"

        if Image and blob:
            with Image.open(BytesIO(blob)) as img:
                result["pixel_width"] = img.width
                result["pixel_height"] = img.height

        return result
    except (AttributeError, KeyError, OSError):
        return None


def extract_slide(slide, slide_index: int, output_dir: Path | None, image_counter: list[int]) -> dict:
    """Extract all content from a single slide."""
    elements = []

    for shape in slide.shapes:
        if shape.has_text_frame:
            paragraphs = extract_text_frame(shape.text_frame)
            if paragraphs:
                role = classify_element(paragraphs)
                elements.append({
                    "type": "text",
                    "role": role,
                    "paragraphs": paragraphs,
                    "position": {
                        "left": shape.left,
                        "top": shape.top,
                        "width": shape.width,
                        "height": shape.height,
                    },
                })

        if shape.has_table:
            table = shape.table
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            elements.append({
                "type": "table",
                "rows": rows,
                "position": {
                    "left": shape.left,
                    "top": shape.top,
                },
            })

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img = extract_image(shape, output_dir, image_id=image_counter[0])
            if img:
                image_counter[0] += 1
                elements.append(img)

    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes = slide.notes_slide.notes_text_frame.text.strip()

    layout_name = ""
    if slide.slide_layout:
        layout_name = slide.slide_layout.name or ""

    return {
        "slide_number": slide_index + 1,
        "layout": layout_name,
        "elements": elements,
        "notes": notes,
    }


def extract_presentation(pptx_path: str, output_dir: Path | None = None) -> dict:
    """Extract full presentation content."""
    prs = Presentation(pptx_path)

    slide_width = prs.slide_width
    slide_height = prs.slide_height

    image_counter = [0]
    slides = []
    for i, slide in enumerate(prs.slides):
        slides.append(extract_slide(slide, i, output_dir, image_counter))

    return {
        "source": str(pptx_path),
        "slide_count": len(slides),
        "dimensions": {
            "width": slide_width,
            "height": slide_height,
            "aspect_ratio": round(slide_width / slide_height, 3) if slide_height else None,
        },
        "slides": slides,
    }


def main():
    parser = argparse.ArgumentParser(description="Extract PPTX content for slide-sage")
    parser.add_argument("input", help="Path to .pptx file")
    parser.add_argument("--output", "-o", help="Output JSON path (default: stdout)")
    parser.add_argument(
        "--images-dir",
        help="Directory to save extracted images (default: embed as data URI)",
    )
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"Error: File not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    output_dir = Path(args.images_dir) if args.images_dir else None
    if output_dir:
        output_dir.mkdir(parents=True, exist_ok=True)

    result = extract_presentation(str(input_path), output_dir)

    json_output = json.dumps(result, indent=2, default=str)

    if args.output:
        Path(args.output).write_text(json_output)
        print(f"Extracted {result['slide_count']} slides to {args.output}")
    else:
        print(json_output)


if __name__ == "__main__":
    main()
